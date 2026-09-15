"""
NewBusiness ETL — win-likelihood forecasting for first-time (new-business) quotes.

ISOLATION: this project's source of truth is the "RSD Scorecard Export"
workbook (required) plus, if present, "External Market Pricing_All Time"
(optional, supplemental — see _load_market_pricing) — both live directly in
NewBusiness/ (one level up from this file). Neither is ever read from, or
written into, the sibling renewal project (ForecastEngine/backend,
ForecastEngine/All-Data, etc.) — see NewBusiness/README.md. This is a
deliberate, separate project: different unit of analysis (a quote, not a
renewal), different label (won/lost sale vs. renewed/not-renewed), different
leakage rules.

Produces:
  data/nb_history.csv   — decided opportunities (Closed Won / Closed Lost)
  data/nb_pipeline.csv  — still-open opportunities (every other Stage)

Run:  python etl_nb.py   (from NewBusiness/backend/)
"""
from __future__ import annotations

import io
import re
import sys
from datetime import datetime
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent))
from app import paths  # noqa: E402

ROOT = paths.export_search_dir()   # NewBusiness/ in dev; per-user state dir when frozen
DATA = paths.backend_dir() / "data"

SHEET = "RSD Scorecard Export"

DECIDED_STAGES = {"Closed Won", "Closed Lost"}

# Sandbox/dummy rows show up in every export under a new name each time
# ("Test Adobe", "TEST HP OPP", "Crumdale Test - BUCA", "Temple ISD - TEST
# CASE - Internal Only", ...) so an exact-name allowlist keeps missing new
# ones -- three of these were sitting in nb_history.csv marked Closed Won,
# inflating the win count and corrupting the rsd/broker target encoding for
# whichever dummy broker they used, and one ("TEST Stop Loss") was live in
# the open pipeline scoring as a top-ranked prospect. Catch the whole family
# by the standalone word "test" in the name instead (word-boundary, not
# substring, so real names like "Whitestone REIT" or "Gastrointestinal
# Associates" aren't affected). The one real company this does flag --
# "Hy-Test Safety Shoe Service" -- is allowlisted back in explicitly.
DUMMY_NAME_RE = re.compile(r"\btest\b", re.IGNORECASE)
REAL_NAME_OVERRIDES = {"Hy-Test Safety Shoe Service"}


def _is_dummy_name(name) -> bool:
    if not isinstance(name, str) or name in REAL_NAME_OVERRIDES:
        return False
    return bool(DUMMY_NAME_RE.search(name))


def _find_sources() -> list[Path]:
    """Every RSD Scorecard Export in ROOT, OLDEST FIRST — not just the newest.

    Why all of them rather than the latest: these exports are not all the same
    pull. Some are the full all-time history; others are narrower slices that
    Salesforce filtered by Effective Date (the September 2026 export carries
    only Effective Date >= 2026-01-01 — 3,340 rows / 81 wins, against the
    all-time August pull's 9,736 rows / 452 wins). Taking `files[-1]` by mtime,
    as this used to, meant dropping in a narrow-but-fresher export SILENTLY
    truncated training history to whatever window that one file happened to
    cover, taking ~370 wins with it. Ordered oldest-first here so
    _merge_exports() can let each newer file correct the ones before it (see
    its docstring) without any file ever being able to delete rows.
    """
    files = sorted(ROOT.glob("RSD Scorecard Export*.xls*"), key=lambda p: p.stat().st_mtime)
    if not files:
        raise FileNotFoundError(
            f"No 'RSD Scorecard Export*.xlsx' found directly in {ROOT} — this "
            "project reads ONLY from those files, nowhere else (see README). "
            "In a packaged build, drop a fresh export into that folder by hand "
            "(the workbook itself is never bundled into the exe)."
        )
    return files


def detect_export_kind(raw_bytes: bytes) -> str | None:
    """Classify an uploaded workbook by its SHEET NAMES, not its filename.

    The RSD Scorecard Export always ships with a sheet literally named
    SHEET regardless of what the file itself gets renamed to -- Salesforce's
    own auto-generated name, a manual "- September" rename, a timestamp
    suffix, whatever. Several files already sitting in NewBusiness/ prove
    the point: same sheet, three different filename shapes. Detecting by
    sheet name instead of a filename prefix means the in-app upload doesn't
    care what the file is called, only what it actually is.

    Returns "scorecard", "pricing", or None if neither expected sheet is
    present (i.e. this isn't a file this project reads at all).
    """
    wb = openpyxl.load_workbook(io.BytesIO(raw_bytes), read_only=True)
    try:
        names = set(wb.sheetnames)
    finally:
        wb.close()
    if SHEET in names:
        return "scorecard"
    if PRICING_SHEET in names:
        return "pricing"
    return None


# The identity of an opportunity ACROSS exports. Deliberately the same subset
# the single-file dedupe below already treats as "the same opportunity" -- one
# notion of row identity in this ETL, not two. Verified unique (zero collisions)
# on both the all-time August and the filtered September exports after
# _normalize_export()'s filters, and it matches 93% of September's rows onto an
# August row; Stage is excluded precisely because it's the field we expect a
# newer export to CHANGE.
MERGE_KEY = ["Opportunity Name", "Created Date", "Associated Broker Account"]


def _merge_key(df: pd.DataFrame) -> pd.Series:
    parts = []
    for name in MERGE_KEY:
        col = _col(df, name)
        if "Date" in name:
            parts.append(pd.to_datetime(df[col], errors="coerce").dt.strftime("%Y-%m-%d").fillna(""))
        else:
            parts.append(df[col].astype(str).str.strip().str.lower())
    return parts[0].str.cat(parts[1:], sep="|")


def _normalize_export(raw: pd.DataFrame) -> pd.DataFrame:
    """Per-file filtering that has to happen BEFORE any cross-file merge, so
    every file is reduced to real, deduplicated New Business rows first and the
    merge only ever compares like with like."""
    raw = raw[raw[_col(raw, "Sales Type")] == "New Business"].copy()
    c_name = _col(raw, "Opportunity Name")
    raw = raw[~raw[c_name].apply(_is_dummy_name)].copy()
    # Full-history exports carry exact duplicate rows for a handful of
    # opportunities (same name/created date/broker repeated verbatim, e.g.
    # a re-run report row) -- keep the first and drop the rest so those
    # opportunities aren't double-weighted in training.
    return raw.drop_duplicates(
        subset=[c_name, _col(raw, "Created Date"), _col(raw, "Associated Broker Account")],
        keep="first",
    ).copy()


def _merge_exports(frames: list[tuple[str, pd.DataFrame]]) -> pd.DataFrame:
    """Fold exports (oldest first) into one frame: newer files CORRECT older
    rows and ADD unseen ones, but can never remove one.

    Union, not replace. A narrow export covering only Effective Date >= 2026
    has nothing to say about a 2024 opportunity -- its silence is a filter
    artifact, not a deletion -- so rows it doesn't mention are carried through
    untouched. That's what keeps the 452-win all-time history intact while
    still picking up September's fresher stages.

    The one subtlety is blanks, where "absent column" and "empty cell" have to
    be told apart or the merge quietly corrupts outcomes:

      * column MISSING from the newer file  -> keep the older value. The newer
        export has no opinion (September drops 'Billing Address Line 1'
        entirely; that address didn't stop existing).
      * column PRESENT but the cell is BLANK -> take the blank. The newer
        export's opinion IS "empty", and overriding it would resurrect stale
        data: 5 rows flipped Closed Lost -> Closed Won in September, which
        clears their Loss Reason. Filling that back from August would leave a
        won deal carrying a loss reason.
    """
    label, base = frames[0]
    base = base.copy()
    base.index = _merge_key(base)
    print(f"  base:    {label} — {len(base)} row(s)")

    for label, nxt in frames[1:]:
        nxt = nxt.copy()
        nxt.index = _merge_key(nxt)
        shared = base.index.intersection(nxt.index)
        added = nxt.index.difference(base.index)

        # Only columns this file actually carries may speak for its rows; every
        # other column keeps whatever the older export had.
        updates = nxt.loc[shared, [c for c in nxt.columns if c in base.columns]]
        changed = int((base.loc[shared, "Stage"].astype(str)
                       != nxt.loc[shared, "Stage"].astype(str)).sum())
        base.loc[shared, updates.columns] = updates

        base = pd.concat([base, nxt.loc[added]], axis=0)
        print(f"  overlay: {label} — {len(shared)} row(s) refreshed "
              f"({changed} stage change(s)), {len(added)} new, "
              f"{len(base) - len(shared) - len(added)} untouched")

    return base.reset_index(drop=True)


PRICING_SHEET = "Market Quotes - All Time"


def _find_pricing_source() -> Path | None:
    """Optional, unlike _find_source(): the model works fine without this
    file (that's how it's worked until now), just with sparser current_max_cost
    / current_renewal / illustrative_max_cost / firm_max_cost coverage."""
    files = sorted(ROOT.glob("External Market Pricing*.xls*"), key=lambda p: p.stat().st_mtime)
    return files[-1] if files else None


def _load_market_pricing() -> pd.DataFrame | None:
    """New Logic's market-pricing waterfall (given verbatim, applied literally):
      - If a Firm quote (Q) exists, use it; otherwise use the Illustrative
        quote (P).
      - If a bound/Written number (W) exists and is CHEAPER than whichever of
        P/Q was picked, use W instead.
      - Compare that one chosen number to the client's Current Max Cost /
        Current Renewal (L).

    Unlike the RSD Scorecard Export (one row per opportunity), this workbook
    carries one row per STOP-LOSS MARKET quoted on an opportunity -- so a
    single client can have several rows, each with its own P/Q/W. New Logic's
    rule is applied per market-quote row first, then the CHEAPEST resulting
    number across that client's markets is kept -- the number a broker would
    actually bring back to the client. current_max_cost/current_renewal
    should be identical across every market row for the same client; taking
    the min (pandas' min skips NaN) recovers the real value wherever at least
    one row has it, which is exactly the "fill the gap" this file is for.

    New Business only (per New Logic's spec) -- Renewal Business rows are
    dropped here so they can never leak into this project, matching the
    ISOLATION rule above.

    Returns None if the workbook isn't present -- this source is optional.
    """
    path = _find_pricing_source()
    if path is None:
        return None
    return _load_market_pricing_from(path)


def _load_market_pricing_from(source) -> pd.DataFrame | None:
    """The actual parse, factored out of _load_market_pricing() so the
    in-app upload preview can run it against uploaded bytes (an io.BytesIO)
    before the file has been saved anywhere -- same logic either way."""
    raw = pd.read_excel(source, sheet_name=PRICING_SHEET, header=0)
    raw = raw[raw["Sales Type"] == "New Business"].copy()
    if raw.empty:
        return None

    L_cur = _bounded(_num(raw, "Current Max Cost"), 1_000, 100_000_000)
    L_ren = _bounded(_num(raw, "Current Renewal"), 1_000, 100_000_000)
    P = _bounded(_num(raw, "Crumdale Max Cost Illustrative"), 1_000, 100_000_000)
    Q = _bounded(_num(raw, "Crumdale Max Cost Firm"), 1_000, 100_000_000)
    W = _bounded(_num(raw, "Total Max Cost"), 1_000, 100_000_000)

    selected = Q.where(Q.notna(), P)
    cheaper_written = W.notna() & (selected.isna() | (W < selected))
    selected = selected.where(~cheaper_written, W)

    mp = pd.DataFrame({
        "k": raw["Opportunity Name"].map(key),
        "mp_current_max_cost": L_cur,
        "mp_current_renewal": L_ren,
        "mp_illustrative_max_cost": P,
        "mp_firm_max_cost": Q,
        "mp_selected_max_cost": selected,
    })
    mp = mp[mp["k"] != ""]
    if mp.empty:
        return None
    return mp.groupby("k").min().reset_index()


def _col(df: pd.DataFrame, *needles: str) -> str:
    """Find a column by substring match, case-insensitive — robust to the
    workbook's en-dash/encoding quirks in a couple of header names. An exact
    (case-insensitive) match wins outright over other substring hits (e.g.
    "Effective Date" vs. the also-present "Effective Date C")."""
    exact = [c for c in df.columns if len(needles) == 1 and c.lower() == needles[0].lower()]
    if len(exact) == 1:
        return exact[0]
    hits = [c for c in df.columns if all(n.lower() in c.lower() for n in needles)]
    if len(hits) != 1:
        raise KeyError(f"expected exactly one column matching {needles}, got {hits}")
    return hits[0]


def _num(df: pd.DataFrame, col: str) -> pd.Series:
    return pd.to_numeric(df[col], errors="coerce")


def _bounded(s: pd.Series, lo: float, hi: float) -> pd.Series:
    """Sanity-bound a numeric column — values outside [lo, hi] are almost
    always export/formula artifacts (Salesforce rollup defaults, div-by-a-
    near-zero-number), and a wrong number is worse than a blank one."""
    return s.where((s >= lo) & (s <= hi))


def key(name) -> str:
    """Name normalization for this project only — same spirit as the renewal
    project's key(), but its own independent copy (no shared code)."""
    if not isinstance(name, str):
        return ""
    s = name.lower()
    s = re.sub(r"[,.\-&'\"]", " ", s)
    s = re.sub(r"\b(inc|llc|corp|co|ltd|company)\b", "", s)
    s = re.sub(r"\s+", " ", s).strip()
    return s


def _days(later: pd.Series, earlier: pd.Series) -> pd.Series:
    d = (pd.to_datetime(later, errors="coerce") - pd.to_datetime(earlier, errors="coerce")).dt.days
    return d.where(d >= 0)  # a negative gap is a data-entry error, not a real lead time


def _row_keys(frame: pd.DataFrame) -> pd.Series:
    """An opportunity's identity in the OUTPUT frames (nb_history.csv /
    nb_pipeline.csv), as opposed to MERGE_KEY's identity in the raw export.
    Shared by _preserve_local_decisions (matching a local edit back to its
    opportunity) and the upload-preview diff report (deciding whether a row
    is new)."""
    return (frame["group_name"].astype(str).str.strip() + "||"
            + pd.to_datetime(frame["created_date"], errors="coerce").dt.strftime("%Y-%m-%d").fillna(""))


LOCAL_SOURCES = ("manual_transfer", "auto_expired")


def _preserve_local_decisions(history: pd.DataFrame) -> tuple[pd.DataFrame, set[str]]:
    """Re-append decisions that were recorded IN THE APP, not in the export.

    Two mechanisms create history rows the workbook doesn't know about yet:
    a manual stage edit to Closed Won/Lost (nb_mode.update_stage) and the
    past-effective-date auto-expire (nb_mode.auto_expire_lost). Without this,
    every ETL rebuild -- including the one /api/retrain runs -- silently threw
    those rows away, because build() reconstructs nb_history.csv from the
    workbook alone. Same guarantee the renewal project's etl_real.py already
    makes for its manual_transfer rows.

    Rows are keyed by (group_name, created_date): if the export has since
    caught up and now carries its own decided row for that opportunity, the
    EXPORT wins (it's the authoritative record, and it has the outcome detail
    fields the local row deliberately left blank) and the local copy is
    dropped rather than duplicated.

    Legacy files written before the `source` column existed are handled by
    inference: any decided row in the old file that the current export has no
    decided row for must have come from one of the two local mechanisms.

    Returns (history, preserved_keys). The caller MUST drop preserved_keys from
    the open pipeline: the export still lists those opportunities as open (it
    predates the in-app decision), so without that the same opportunity would
    sit in nb_history.csv AND nb_pipeline.csv at once -- trained on as decided
    and scored as open, and double-counted in every pipeline total.
    """
    path = DATA / "nb_history.csv"
    if not path.exists():
        return history, set()

    prev = pd.read_csv(path)
    if prev.empty:
        return history, set()

    if "source" in prev.columns:
        local = prev[prev["source"].isin(LOCAL_SOURCES)].copy()
    else:
        local = prev[~_row_keys(prev).isin(set(_row_keys(history)))].copy()
        local["source"] = "manual_transfer"   # can't tell which mechanism, retroactively
    if local.empty:
        return history, set()

    local = local[~_row_keys(local).isin(set(_row_keys(history)))]
    if local.empty:
        print("  preserved local decisions: 0 (export has caught up on all of them)")
        return history, set()

    local = local.reindex(columns=history.columns).astype(
        {c: t for c, t in history.dtypes.items() if t == object})
    # `history` at this point holds real datetime64 Timestamps for eff_date/
    # created_date/notice_of_sale_date (freshly parsed from the export by
    # build()); `local` holds whatever pd.read_csv gave back for those columns
    # -- plain strings, since the on-disk file has no dtype info. Concatenating
    # a Timestamp column with a string column silently downgrades the WHOLE
    # merged column to dtype=object, and pandas's CSV writer then falls back to
    # str(x) per cell instead of its clean-date formatting for a pure
    # datetime64 column -- str(Timestamp) always includes " 00:00:00" even at
    # midnight. That bug doesn't just affect the few preserved rows; it
    # corrupts the date formatting for every OTHER row in the file the moment
    # the column stops being uniformly datetime64. Re-parsing both sides to a
    # real Timestamp before concatenating keeps the column uniform and the
    # written dates clean (caught 2026-08-26: 9,394 of 9,396 history rows had
    # picked up the " 00:00:00" suffix and were rendering as "Invalid Date" in
    # the Win/Loss Database).
    for col in ("eff_date", "created_date", "notice_of_sale_date"):
        if col in local.columns:
            local[col] = pd.to_datetime(local[col], errors="coerce")
    merged = pd.concat([history, local], ignore_index=True)
    print(f"  preserved local decisions: {len(local)} row(s) recorded in-app, not in the export")
    return merged, set(_row_keys(local))


def build(extra_upload: tuple[str, bytes] | None = None, write: bool = True) -> dict:
    """Run the whole ETL: merge every export in ROOT, oldest first, into
    history + pipeline.

    extra_upload -- (label, raw .xlsx bytes) for a file that ISN'T on disk
    yet, folded into the merge as the newest source (in-memory, via
    io.BytesIO — never touches ROOT). This is what powers the in-app upload
    preview: a candidate file can be merged and reported on without saving
    it or committing anything, exactly the same code path a file dropped
    into the folder by hand would take. See app/main.py's
    /api/upload/scorecard/preview.

    write=False skips the CSV write (and the print of what got written) --
    also for preview, so a rejected upload can never have touched
    nb_history.csv / nb_pipeline.csv.
    """
    paths_in = _find_sources()
    sources = [(p.name, _normalize_export(pd.read_excel(p, sheet_name=SHEET, header=0)))
               for p in paths_in]
    source_names = [p.name for p in paths_in]
    if extra_upload is not None:
        label, raw_bytes = extra_upload
        sources.append((label, _normalize_export(
            pd.read_excel(io.BytesIO(raw_bytes), sheet_name=SHEET, header=0))))
        source_names.append(label)
    print(f"merging {len(sources)} export(s) from {ROOT}"
          + (" (+ 1 pending upload, not yet saved)" if extra_upload else "") + ":")
    raw = _merge_exports(sources)

    c_name = _col(raw, "Opportunity Name")
    c_rsd = _col(raw, "Primary Owner")
    c_uw = _col(raw, "Underwriter Owner")
    c_broker = _col(raw, "Associated Broker Account")
    c_product = _col(raw, "Type of Quote Requested")
    c_created = _col(raw, "Created Date")
    c_eff = _col(raw, "Effective Date")
    c_stage = _col(raw, "Stage")
    c_lives = _col(raw, "No. of Enrolled Employees")
    c_cur_max = _col(raw, "Current Max Cost")
    c_illus_max = _col(raw, "Crumdale Max Cost", "Illustrative")
    c_firm_max = _col(raw, "Crumdale Max Cost", "Firm")
    c_cur_renewal = _col(raw, "Current Renewal")
    c_dt_illus = _col(raw, "Date Illustrative Quote sent to RSD")
    c_dt_uw_firm = _col(raw, "Date UW Completed Firm Quote")
    c_dt_firm = _col(raw, "Date Firm Quote sent to RSD")
    c_notice = _col(raw, "Notice of Sale Date")
    c_loss_reason = _col(raw, "Loss Reason")
    c_loss_other = _col(raw, "Loss Reason - Other")
    c_isl = _col(raw, "ISL Deductible")
    c_laser1, c_laser2, c_laser3, c_laser4 = (
        _col(raw, "Laser 1"), _col(raw, "Laser 2"), _col(raw, "Laser 3"), _col(raw, "Laser 4"))
    c_city = _col(raw, "Billing City")
    c_state = _col(raw, "Billing State")
    c_zip = _col(raw, "Billing Zip")
    c_industry0 = _col(raw, "Industry")
    c_industry1 = _col(raw, "Industry.1")

    out = pd.DataFrame()
    out["group_name"] = raw[c_name].astype(str).str.strip()
    out["k"] = out["group_name"].map(key)
    out["rsd"] = raw[c_rsd].astype(str).str.strip().replace({"nan": np.nan})
    out["underwriter"] = raw[c_uw].astype(str).str.strip().replace({"nan": np.nan})
    out["broker"] = raw[c_broker].astype(str).str.strip().replace({"nan": np.nan})
    out["product"] = raw[c_product].astype(str).str.strip().replace({"nan": np.nan})
    # Industry.1 first, falling back to Industry. Industry.1 is the better
    # PRIMARY choice because it's the one populated where it matters most --
    # 93.6% on the open pipeline we actually have to score, vs. 72.1% for
    # Industry. But across the whole export Industry is the fuller column
    # (90.8% vs 87.5%), so coalescing rather than picking one outright lifts
    # combined coverage to ~97% and strictly dominates either alone.
    # (Industry.2 is dropped: it duplicates Industry.1 almost everywhere it
    # isn't blank, so it adds no independent coverage.)
    _ind1 = raw[c_industry1].astype(str).str.strip().replace({"nan": np.nan, "": np.nan})
    _ind0 = raw[c_industry0].astype(str).str.strip().replace({"nan": np.nan, "": np.nan})
    out["industry"] = _ind1.fillna(_ind0)
    out["billing_city"] = raw[c_city].astype(str).str.strip().replace({"nan": np.nan})
    out["billing_state"] = raw[c_state].astype(str).str.strip().replace({"nan": np.nan})
    # Zip-3 prefix, not the full zip: the 5-digit code has 2,772 distinct
    # values against 446 wins (pure overfitting bait even target-encoded),
    # while the 3-digit prefix is a real regional grouping (~sectional center
    # facility) with enough quotes per bucket to carry a stable win rate.
    out["billing_zip3"] = (raw[c_zip].astype(str).str.strip().str.extract(r"^(\d{3})")[0]
                            .replace({"nan": np.nan, "": np.nan}))

    out["created_date"] = pd.to_datetime(raw[c_created], errors="coerce")
    out["eff_date"] = pd.to_datetime(raw[c_eff], errors="coerce")
    out["eff_month"] = out["eff_date"].dt.to_period("M").astype(str)

    out["stage"] = raw[c_stage].astype(str).str.strip()
    out["lives"] = _bounded(_num(raw, c_lives), 1, 20000)

    # Pricing fields are the messiest in this export (Salesforce rollup
    # placeholders show up as literal 0/1 instead of blank, and a few rows have
    # absurd multi-billion-dollar values from a formula dividing by ~0). Bound
    # every dollar figure to a plausible annual-premium range and every ratio to
    # a plausible -95%..+500% band; outside that, leave it blank rather than
    # feed the model (or the dashboard) a number that's obviously wrong.
    out["current_max_cost"] = _bounded(_num(raw, c_cur_max), 1_000, 100_000_000)
    out["illustrative_max_cost"] = _bounded(_num(raw, c_illus_max), 1_000, 100_000_000)
    out["current_renewal"] = _bounded(_num(raw, c_cur_renewal), 1_000, 100_000_000)
    out["firm_max_cost"] = _bounded(_num(raw, c_firm_max), 1_000, 100_000_000)

    # --- Market-pricing gap-fill (External Market Pricing_All Time workbook,
    # if dropped into NewBusiness/ -- see _load_market_pricing) --- Fills GAPS
    # ONLY (never overwrites a real value already on hand) in the four cost
    # fields above -- this is task 1, nothing more: no new column, just fewer
    # blanks in the fields the model already uses.
    mp = _load_market_pricing()
    _mp_selected = pd.Series(np.nan, index=out.index)   # task-2 input, not a saved column -- see below
    if mp is not None:
        out = out.merge(mp, on="k", how="left")
        for col, mp_col in (
            ("current_max_cost", "mp_current_max_cost"),
            ("current_renewal", "mp_current_renewal"),
            ("illustrative_max_cost", "mp_illustrative_max_cost"),
            ("firm_max_cost", "mp_firm_max_cost"),
        ):
            fill = out[col].isna() & out[mp_col].notna()
            out.loc[fill, col] = out.loc[fill, mp_col]
        _mp_selected = out["mp_selected_max_cost"]
        out = out.drop(columns=["mp_current_max_cost", "mp_current_renewal",
                                 "mp_illustrative_max_cost", "mp_firm_max_cost",
                                 "mp_selected_max_cost"])

    # Computed directly from the two cost fields rather than read off
    # Salesforce's own "% vs current" formula column -- verified to match it
    # exactly on every row that has one (1,644/1,644), and computing it
    # ourselves means it's still available on exports (e.g. the fuller
    # all-time pull) that don't carry that report-layer column at all.
    # Recomputed AFTER the backfill above so it benefits from any newly-filled
    # current_max_cost/illustrative_max_cost -- unchanged logic otherwise.
    out["pct_vs_current"] = _bounded(out["illustrative_max_cost"] / out["current_max_cost"] - 1, -0.95, 5.0)

    # --- Task 2: New Logic's comparison rule, verbatim, nothing extra ---
    # "If P and no Q, compare P to L. If Q, compare Q to L. If W and W is less
    # than Q (if Q) or P (if no Q), compare W to L." _mp_selected (computed
    # above in _load_market_pricing, per market-quote row then cheapest-
    # across-markets) already IS that chosen number wherever a market quote
    # exists; falls back to firm-then-illustrative when it doesn't, so the
    # comparison still runs off the RSD export alone. The two lines below are
    # the literal "compare ... to L" step, against both of Crumdale's L's
    # (Current Max Cost, Current Renewal) -- these are the only two new
    # columns this task produces; the selected number itself is an
    # intermediate and isn't saved.
    _selected = _mp_selected.where(
        _mp_selected.notna(),
        out["firm_max_cost"].where(out["firm_max_cost"].notna(), out["illustrative_max_cost"]))
    out["pct_vs_current_full"] = _bounded(_selected / out["current_max_cost"] - 1, -0.95, 5.0)
    out["pct_vs_renewal"] = _bounded(_selected / out["current_renewal"] - 1, -0.95, 5.0)

    # Range vs Current / Renewal Range / Firm w Laser Liability were dropped
    # entirely (not just left blank) -- they are Salesforce report-layer
    # formula columns that this export (the fuller all-time RSD Scorecard
    # pull) never carries at all, on EITHER decided history or the open
    # pipeline (confirmed against the raw column headers directly, not
    # inferred from blank cells). Keeping a column that is 0% populated on
    # both sides isn't an honest "missing value" case -- there's no signal to
    # honestly represent, so it's removed from the data model rather than
    # carried as permanent dead weight (see also NUMERIC/ONEHOT_CATEGORICAL in
    # train_nb.py and data_quality_nb.py's FIELD_GROUPS).
    #
    # RUS:DTQ? and Competitive? are dropped for the same reason -- neither
    # column exists in this export either. An earlier version of this ETL
    # defaulted them to a constant 0 when the column was missing, which
    # produced zero-variance "flags" that read as 100% complete on the Data
    # Quality page (misleadingly -- a constant isn't real completeness) and
    # could never inform the model or the rule-based reason codes in
    # insights_nb.py (drivers() no longer references either field).
    out["isl_deductible"] = _bounded(_num(raw, c_isl), 1_000, 5_000_000)

    laser_cols = [_num(raw, c) for c in (c_laser1, c_laser2, c_laser3, c_laser4)]
    lasers = pd.concat(laser_cols, axis=1)
    out["laser_liability"] = lasers.sum(axis=1, min_count=1)
    # How many lasers, not just their combined dollar value -- a distinct signal
    # from the total (one $2M laser reads very differently than four $500k ones).
    out["laser_count"] = lasers.notna().sum(axis=1)

    dt_illus = pd.to_datetime(raw[c_dt_illus], errors="coerce")
    dt_uw_firm = pd.to_datetime(raw[c_dt_uw_firm], errors="coerce")
    dt_firm = pd.to_datetime(raw[c_dt_firm], errors="coerce")

    # Whether a real quote number/date exists -- deliberately NOT falling back
    # to stage membership. An earlier version OR'd in "stage in {...Closed Won,
    # Closed Lost}", which meant every row in nb_history.csv (all of them
    # decided, i.e. always a member of that set) read as 1 regardless of
    # whether a quote was ever actually produced -- zero variance, so the
    # model could never learn anything from it (confirmed at 0.0% importance).
    # Checking the underlying signal directly gives real, varying values on
    # both history and the open pipeline.
    out["has_illustrative_quote"] = (out["illustrative_max_cost"].notna() | dt_illus.notna()).astype(int)
    out["has_firm_quote"] = (
        out["firm_max_cost"].notna() | dt_uw_firm.notna() | dt_firm.notna()
    ).astype(int)

    out["days_created_to_eff"] = _days(out["eff_date"], out["created_date"])
    out["days_created_to_illustrative"] = _days(dt_illus, out["created_date"])
    out["days_illustrative_to_firm"] = _days(dt_firm, dt_illus)
    # UW turnaround, broken into its two legs: how long underwriting itself took
    # once the illustrative quote was out, and how long the finished firm quote
    # sat before it actually reached the RSD -- a slow internal handoff is a
    # distinct (and controllable) risk from slow underwriting itself.
    out["days_illustrative_to_uw_complete"] = _days(dt_uw_firm, dt_illus)
    out["days_uw_complete_to_firm_sent"] = _days(dt_firm, dt_uw_firm)

    # Per-life (per-enrolled-employee) normalized cost -- the raw dollar figures
    # mostly just proxy for group size, which we already have via `lives`;
    # normalizing gives a size-independent read on how rich/lean the quote is.
    for src, dst in (("current_max_cost", "current_cost_per_life"),
                      ("illustrative_max_cost", "illustrative_cost_per_life"),
                      ("firm_max_cost", "firm_cost_per_life")):
        out[dst] = _bounded(out[src] / out["lives"].where(out["lives"] > 0), 200, 60_000)

    # --- outcome / leakage-sensitive fields ---
    # Notice of Sale Date is present in 98.7% of wins and 0.1% of losses/open
    # rows — it IS the win confirmation date, not a predictive signal. Loss
    # Reason is assigned only after a Closed Lost decision. Neither is ever a
    # model feature; both are kept in nb_history.csv purely for the historical
    # database view.
    out["notice_of_sale_date"] = pd.to_datetime(raw[c_notice], errors="coerce")
    out["loss_reason"] = raw[c_loss_reason].astype(str).str.strip().replace({"nan": np.nan})
    out["loss_reason_other"] = raw[c_loss_other].astype(str).str.strip().replace({"nan": np.nan})
    # Computed from Stage rather than read off the "Sold" formula column --
    # verified to match it exactly on every historical row (3,051/3,051) --
    # for the same reason as pct_vs_current above: it's still available on
    # exports that don't carry that report-layer column.
    out["sold"] = (out["stage"] == "Closed Won").astype(int)
    out["decided"] = out["stage"].isin(DECIDED_STAGES)

    out["source"] = "export"

    history = out[out["decided"]].drop(columns=["decided"]).reset_index(drop=True)
    pipeline = out[~out["decided"]].drop(columns=["decided", "sold", "notice_of_sale_date",
                                                    "loss_reason", "loss_reason_other"]).reset_index(drop=True)

    history, local_keys = _preserve_local_decisions(history)
    if local_keys:
        pk = (pipeline["group_name"].astype(str).str.strip() + "||"
              + pipeline["created_date"].dt.strftime("%Y-%m-%d").fillna(""))
        before = len(pipeline)
        pipeline = pipeline[~pk.isin(local_keys)].reset_index(drop=True)
        if before != len(pipeline):
            print(f"  removed {before - len(pipeline)} row(s) from the open pipeline "
                  f"— already decided in-app (kept in history, not double-counted)")

    # Belt-and-suspenders: force every date column to a real Timestamp right
    # before writing, whatever dtype it arrived in. A single non-Timestamp
    # value anywhere in one of these columns (a stray string from a merge, a
    # NaT-as-object edge case) downgrades the WHOLE column to dtype=object,
    # which silently breaks pandas's clean-date CSV formatting for every OTHER
    # row too (see _preserve_local_decisions's docstring for exactly how that
    # happened here on 2026-08-26). This is the last line of defense, not a
    # substitute for keeping dtypes clean upstream.
    for df in (history, pipeline):
        for col in ("eff_date", "created_date", "notice_of_sale_date"):
            if col in df.columns:
                df[col] = pd.to_datetime(df[col], errors="coerce")

    if write:
        DATA.mkdir(parents=True, exist_ok=True)
        history.to_csv(DATA / "nb_history.csv", index=False)
        pipeline.to_csv(DATA / "nb_pipeline.csv", index=False)

    print(f"sources: {', '.join(source_names)}")
    print(f"New Business rows: {len(raw)}")
    print(f"  decided (history): {len(history)}  (won {int(history['sold'].sum())} / "
          f"lost {int((history['sold'] == 0).sum())})")
    print(f"  open (pipeline):   {len(pipeline)}")
    print(f"win rate (decided only): {history['sold'].mean():.1%}")
    return {"history": history, "pipeline": pipeline, "sources": source_names, "raw_rows": len(raw)}


# ---------------------------------------------------------------------------
# In-app upload: preview a candidate file before it touches disk, then apply
# it for real. See app/main.py's /api/upload/scorecard/preview and /apply.
#
# Why a preview step at all, given the merge logic above already only ever
# adds/corrects rows and can't delete history: a narrower export can still
# silently mean "many fewer NEW rows than expected" (the September RSD pull
# arriving pre-filtered to Effective Date >= 2026-01-01 is exactly this,
# see DEPLOYMENT_PLAN.md) -- nothing about the merge being non-destructive
# catches that on its own. Showing the numbers before they go live is the
# same discipline the CLI flow already relies on a human reading the printed
# output for; this just surfaces it in the UI instead of a terminal.
# ---------------------------------------------------------------------------

def _read_csv_or_empty(path: Path, like: pd.DataFrame) -> pd.DataFrame:
    return pd.read_csv(path) if path.exists() else pd.DataFrame(columns=like.columns)


def _diff_report(new_history: pd.DataFrame, new_pipeline: pd.DataFrame) -> dict:
    """Compare a candidate history/pipeline pair against what's currently on
    disk (i.e. currently served) and summarize what would change."""
    old_history = _read_csv_or_empty(DATA / "nb_history.csv", new_history)
    old_pipeline = _read_csv_or_empty(DATA / "nb_pipeline.csv", new_pipeline)

    old_h_keys = set(_row_keys(old_history)) if not old_history.empty else set()
    old_p_keys = set(_row_keys(old_pipeline)) if not old_pipeline.empty else set()
    new_h_keys = _row_keys(new_history)
    new_p_keys = _row_keys(new_pipeline)

    total_before = len(old_history) + len(old_pipeline)
    total_after = len(new_history) + len(new_pipeline)

    warning = None
    if total_before and total_after < total_before * 0.9:
        warning = (f"Row count would drop from {total_before:,} to {total_after:,} "
                   f"({total_before - total_after:,} fewer) -- that looks like a filtered "
                   "or partial export, not a smaller book. Check the source file before applying.")

    return {
        "new_decided_outcomes": int((~new_h_keys.isin(old_h_keys)).sum()),
        "new_open_quotes": int((~new_p_keys.isin(old_p_keys)).sum()),
        "total_history_before": len(old_history), "total_history_after": len(new_history),
        "total_pipeline_before": len(old_pipeline), "total_pipeline_after": len(new_pipeline),
        "warning": warning,
    }


def preview_upload(raw_bytes: bytes, kind: str, filename: str) -> dict:
    """What WOULD happen if `filename` were saved and the ETL rerun --
    computed without writing anything, so a bad file can be safely rejected."""
    if kind == "scorecard":
        result = build(extra_upload=(filename, raw_bytes), write=False)
        report = _diff_report(result["history"], result["pipeline"])
        report["sources"] = result["sources"]
        return report
    if kind == "pricing":
        mp = _load_market_pricing_from(io.BytesIO(raw_bytes))
        return {"opportunities_priced": 0 if mp is None else len(mp)}
    raise ValueError(f"unknown upload kind {kind!r}")


def apply_upload(raw_bytes: bytes, kind: str, filename: str) -> dict:
    """Save `filename` into ROOT (so it's just another export from now on,
    same as one dropped in by hand) and rerun the ETL for real. Does NOT
    retrain or rescore -- app/main.py's apply endpoint runs build_book.py
    right after this, same as every other data-import path in this app."""
    ROOT.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y-%m-%d %H%M%S")
    prefix = "RSD Scorecard Export" if kind == "scorecard" else "External Market Pricing_All Time"
    dest = ROOT / f"{prefix} - uploaded {stamp}.xlsx"
    dest.write_bytes(raw_bytes)

    if kind == "scorecard":
        result = build(write=True)
        return {"saved_as": dest.name, "sources": result["sources"],
                "history_rows": len(result["history"]), "pipeline_rows": len(result["pipeline"])}
    if kind == "pricing":
        return {"saved_as": dest.name}
    raise ValueError(f"unknown upload kind {kind!r}")


if __name__ == "__main__":
    build()
